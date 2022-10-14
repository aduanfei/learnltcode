import requests
import docx
import xlrd
import os
from bs4 import BeautifulSoup
import time

import random

data=xlrd.open_workbook(r"C:\Users\之子言归\Desktop\2022上半年主动宣传\dnw.xlsx")
table=data.sheets()[0]
failed=[]
def get1():
    for i in range(2,51):
        #print(table.cell(i,1),table.cell(i,2))
        #time.sleep(random.randint(1,3))
        ptml=table.cell(i,2).value
        title=table.cell(i,1).value
        try:
            print(ptml)
            content=requests.get(ptml)
            cont=BeautifulSoup(content.text,"lxml")
            messa=cont.find(id="new_message_id").findAll("p")
            ps=title+"\n"
            for p in messa:
                ps=ps+p.text+"\n"
            rpath=r"C:\Users\之子言归\Desktop\2022上半年主动宣传\新闻稿"
            rpath+="\\"+title+".docx"
            if  os.path.exists(rpath):
                os.remove(rpath)
            doc=docx.Document()
            doc.add_paragraph(ps)
            doc.save(rpath)
            print(title+" saved!")
        except Exception:
            failed.append((title,ptml))
    print(failed)

def get2():
    ptml="http://www.xxf315.com/index.php?s=news&c=search&keyword=%E5%85%B4%E4%B8%9A%E9%93%B6%E8%A1%8C&page=2"
    data = xlrd.open_workbook(r"C:\Users\之子言归\Desktop\2022上半年主动宣传\dnw.xlsx")
    table = data.sheets()[0]
    i=0
    content = requests.get(ptml)
    cont = BeautifulSoup(content.text, "lxml")
    titles=cont.findAll("div",class_="search-post-foot")
    for t in titles:
        title=t.find("a")
        #print(title)
        date=t.find("div",class_="search-post-meta").a.text

        print(date,title["title"],title["href"])




get2()